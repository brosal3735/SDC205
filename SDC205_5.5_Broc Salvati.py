# SDC205 Final Performance Assessment
# Student ID: Brosal3735

import os
import csv
from datetime import datetime

import openpyxl
from openpyxl.chart import PieChart, Reference
import matplotlib.pyplot as plt


STUDENT_ID = "Brosal3735"
BASE_FOLDER = r"C:\FinalExam"
CSV_FILE = "final.csv"
XLSX_FILE = "final.xlsx"


def askUser():
    total = 0

    # Loop logic: ask the user for 5 numbers, convert each to int, and keep a running total.
    for _ in range(5):
        num = int(input("Please enter a number: "))
        total += num

    print(f"The sum for the 5 numbers entered is: {total}")


def askIncome():
    csv_path = os.path.join(BASE_FOLDER, CSV_FILE)

    # Loop logic: ask for 5 name/income pairs and append each pair as a new CSV row to final.csv.
    # We open the file in append mode so the 4 existing rows remain and we add 5 more rows.
    with open(csv_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)

        for _ in range(5):
            name = input("Please enter a name: ")
            income = input("Please enter their income: ")
            writer.writerow([name, income])


def _read_csv_data():
    csv_path = os.path.join(BASE_FOLDER, CSV_FILE)
    names = []
    incomes = []

    with open(csv_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row:
                continue
            if len(row) < 2:
                continue

            name = str(row[0]).strip()
            income_str = str(row[1]).strip()

            if name == "" or income_str == "":
                continue

            try:
                income_val = int(float(income_str))
            except ValueError:
                continue

            names.append(name)
            incomes.append(income_val)

    return names, incomes


def excelPie():
    names, incomes = _read_csv_data()
    xlsx_path = os.path.join(BASE_FOLDER, XLSX_FILE)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    ws["A1"] = "Name"
    ws["B1"] = "Income"

    for i in range(len(names)):
        ws.cell(row=i + 2, column=1, value=names[i])
        ws.cell(row=i + 2, column=2, value=incomes[i])

    # Creates the data reference used by the pie chart (Income values only, not the header).
    data = Reference(ws, min_col=2, min_row=2, max_row=len(names) + 1)

    # Creates the label reference used by the pie chart (Name labels only).
    labels = Reference(ws, min_col=1, min_row=2, max_row=len(names) + 1)

    # Creates a PieChart object in openpyxl.
    pie = PieChart()

    # Sets the pie chart title to StudentID plus today's date.
    today_str = datetime.now().strftime("%B %d, %Y")
    pie.title = f"{STUDENT_ID} {today_str}"

    # Adds the income data series to the pie chart.
    pie.add_data(data)

    # Sets the category labels (names) that appear on the pie chart.
    pie.set_categories(labels)

    # Adds the pie chart to the worksheet with its top-left corner at cell A11.
    ws.add_chart(pie, "A11")

    wb.save(xlsx_path)


def verticalBar():
    names, incomes = _read_csv_data()

    today_str = datetime.now().strftime("%B %d, %Y")
    plt.figure()
    plt.bar(names, incomes)
    plt.xlabel("Name")
    plt.ylabel("Income")
    plt.title(f"{STUDENT_ID} {today_str}")
    plt.xticks(rotation=0)
    plt.tight_layout()
    plt.show()


def main():
    os.makedirs(BASE_FOLDER, exist_ok=True)
    os.chdir(BASE_FOLDER)

    askUser()
    askIncome()
    excelPie()
    verticalBar()


if __name__ == "__main__":
    main()

